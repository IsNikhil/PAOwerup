"""Reading resumes (PDF / DOCX / TXT) and writing the improved resume as DOCX."""
from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.schemas import ResumeDoc

SUPPORTED = (".pdf", ".docx", ".txt", ".md")


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [page.extract_text(x_tolerance=1.5) or "" for page in pdf.pages]
        text = "\n".join(pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type. Upload one of: {', '.join(SUPPORTED)}")
    text = text.replace("\r", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ---------------------------------------------------------------- DOCX export
ACCENT = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)


def _rule(paragraph) -> None:
    """Thin bottom border under a section heading."""
    p_pr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "CBD5E1")):
        bottom.set(qn(k), v)
    border.append(bottom)
    p_pr.append(border)


def _spacing(paragraph, before: int = 0, after: int = 2) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)


def build_docx(resume: ResumeDoc) -> bytes:
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(54)
        section.top_margin = section.bottom_margin = Pt(46)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    if resume.name:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(resume.name)
        run.bold = True
        run.font.size = Pt(20)
        run.font.color.rgb = ACCENT
        _spacing(p, 0, 0)
    if resume.headline:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(resume.headline)
        run.font.size = Pt(11)
        run.font.color.rgb = MUTED
        _spacing(p, 0, 0)
    if resume.contact:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(resume.contact)
        run.font.size = Pt(9.5)
        run.font.color.rgb = MUTED
        _spacing(p, 0, 6)

    def heading(text: str) -> None:
        p = doc.add_paragraph()
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = ACCENT
        _spacing(p, 10, 3)
        _rule(p)

    if resume.summary:
        heading("Summary")
        _spacing(doc.add_paragraph(resume.summary), 0, 2)

    for section in resume.sections:
        heading(section.title)
        for item in section.items:
            if item.heading or item.dates:
                p = doc.add_paragraph()
                _spacing(p, 4, 0)
                run = p.add_run(item.heading)
                run.bold = True
                if item.subheading:
                    sub = p.add_run(f" — {item.subheading}")
                    sub.font.color.rgb = MUTED
                if item.dates:
                    p.add_run("\t" + item.dates).font.color.rgb = MUTED
                    p.paragraph_format.tab_stops.add_tab_stop(Pt(504), alignment=2)  # right-aligned
            for bullet in item.bullets:
                bp = doc.add_paragraph(bullet, style="List Bullet")
                _spacing(bp, 0, 1)

    if resume.skills:
        heading("Skills")
        _spacing(doc.add_paragraph(", ".join(resume.skills)), 0, 2)

    if resume.skills_to_develop:
        heading("Skills to develop (not yet experience)")
        p = doc.add_paragraph()
        run = p.add_run("Currently learning / planned: " + ", ".join(resume.skills_to_develop))
        run.italic = True
        run.font.color.rgb = MUTED

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
